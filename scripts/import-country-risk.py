from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "src" / "data" / "country-risk"
LATEST_TS = ROOT / "src" / "data" / "countryRiskLatest.ts"
REPORT_DIR = ROOT / "reports"


def clean_number(value: Any) -> float | None:
    if value is None or value in ("NA", "#N/A", ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def iso_date(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def version_from_date(value: str) -> str:
    match = re.match(r"(\d{4})-(\d{2})", value)
    if not match:
        raise ValueError(f"Cannot derive version from update date: {value}")
    return f"{match.group(1)}-{match.group(2)}"


def sheet_rows(wb: Any, sheet_name: str) -> list[tuple[Any, ...]]:
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Workbook is missing required sheet: {sheet_name}")
    return list(wb[sheet_name].iter_rows(values_only=True))


def find_metadata(wb: Any) -> dict[str, Any]:
    rows = sheet_rows(wb, "ERPs by country")
    update_date = None
    mature = None
    us_erp = None
    rel_vol = None

    for row in rows[:10]:
        values = list(row)
        text = " ".join(str(v) for v in values if v is not None).lower()
        if values and values[0] == "Date of update:" and values[1]:
            update_date = iso_date(values[1])
        for value in values:
            if not isinstance(value, (int, float)):
                continue
            if "mature equity market" in text:
                mature = float(value)
            elif "risk premium for the us" in text:
                us_erp = float(value)
            elif "multiplier" in text:
                rel_vol = float(value)

    if update_date is None or mature is None or us_erp is None or rel_vol is None:
        raise ValueError("Could not extract update date, mature ERP, US ERP, and relative volatility metadata")

    return {
        "updateDate": update_date,
        "version": version_from_date(update_date),
        "matureMarketPremium": mature,
        "usErp": us_erp,
        "relativeEquityVolatility": rel_vol,
    }


def extract_countries(wb: Any) -> list[dict[str, Any]]:
    rows = sheet_rows(wb, "Regional breakdown")
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    index = {header: i for i, header in enumerate(headers)}
    required = [
        "Country",
        "GDP (in millions) in 2024",
        "Moody's rating",
        "Sovereign CDS",
        "Adj. Default Spread",
        "Equity Risk Premium",
        "Country Risk Premium",
        "Corporate Tax Rate",
        "Region",
    ]
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(f"Regional breakdown is missing columns: {', '.join(missing)}")

    countries: list[dict[str, Any]] = []
    for row in rows[1:]:
        country = row[index["Country"]] if index["Country"] < len(row) else None
        if not country:
            continue
        erp = clean_number(row[index["Equity Risk Premium"]])
        crp = clean_number(row[index["Country Risk Premium"]])
        spread = clean_number(row[index["Adj. Default Spread"]])
        if erp is None or crp is None or spread is None:
            continue

        rating = row[index["Moody's rating"]]
        countries.append({
            "country": str(country).strip(),
            "region": str(row[index["Region"]] or "").strip(),
            "moodyRating": None if rating in (None, "NA", "#N/A") else str(rating).strip(),
            "sovereignCds": clean_number(row[index["Sovereign CDS"]]),
            "adjustedDefaultSpread": spread,
            "totalEquityRiskPremium": erp,
            "countryRiskPremium": crp,
            "taxRate": clean_number(row[index["Corporate Tax Rate"]]),
            "gdpMillions": clean_number(row[index["GDP (in millions) in 2024"]]),
        })

    if len(countries) < 120:
        raise ValueError(f"Country count looks too low: {len(countries)}")

    return countries


def validate_dataset(dataset: dict[str, Any]) -> list[str]:
    warnings: list[str] = []
    names = [c["country"] for c in dataset["countries"]]
    if len(names) != len(set(names)):
        warnings.append("Duplicate countries detected")

    for country_name in ("United States", "Thailand", "Australia"):
        if country_name not in names:
            raise ValueError(f"Smoke-test country missing: {country_name}")

    mature = dataset["matureMarketPremium"]
    for record in dataset["countries"]:
        if record["countryRiskPremium"] < -0.000001:
            raise ValueError(f"Negative CRP for {record['country']}")
        if record["totalEquityRiskPremium"] + 0.000001 < mature:
            raise ValueError(f"ERP below mature market premium for {record['country']}")
        tax_rate = record.get("taxRate")
        if tax_rate is not None and not 0 <= tax_rate <= 0.6:
            warnings.append(f"Tax rate outside normal range for {record['country']}: {tax_rate}")

    return warnings


def write_latest_ts(dataset: dict[str, Any]) -> None:
    body = json.dumps(dataset, ensure_ascii=False, indent=2)
    LATEST_TS.write_text(
        "export interface CountryRiskRecord {\n"
        "  country: string\n"
        "  region: string\n"
        "  moodyRating?: string | null\n"
        "  sovereignCds?: number | null\n"
        "  adjustedDefaultSpread: number\n"
        "  totalEquityRiskPremium: number\n"
        "  countryRiskPremium: number\n"
        "  taxRate?: number | null\n"
        "  gdpMillions?: number | null\n"
        "}\n\n"
        "export interface CountryRiskDataset {\n"
        "  version: string\n"
        "  sourceFile: string\n"
        "  updateDate: string\n"
        "  matureMarketPremium: number\n"
        "  usErp: number\n"
        "  relativeEquityVolatility: number\n"
        "  countries: CountryRiskRecord[]\n"
        "}\n\n"
        f"const countryRiskLatest: CountryRiskDataset = {body}\n\n"
        "export default countryRiskLatest\n",
        encoding="utf-8",
    )


def write_report(dataset: dict[str, Any], warnings: list[str]) -> None:
    top = sorted(dataset["countries"], key=lambda c: c["countryRiskPremium"], reverse=True)[:10]
    lines = [
        f"# Country Risk Update {dataset['version']}",
        "",
        f"- Source file: `{dataset['sourceFile']}`",
        f"- Update date: `{dataset['updateDate']}`",
        f"- Countries: `{len(dataset['countries'])}`",
        f"- Mature market premium: `{dataset['matureMarketPremium']:.4%}`",
        f"- US ERP: `{dataset['usErp']:.4%}`",
        f"- Relative equity volatility: `{dataset['relativeEquityVolatility']:.4f}`",
        "",
        "## Highest Country Risk Premiums",
        "",
    ]
    for record in top:
        lines.append(f"- {record['country']}: {record['countryRiskPremium']:.2%} CRP, {record['totalEquityRiskPremium']:.2%} ERP")
    if warnings:
        lines.extend(["", "## Warnings", ""])
        lines.extend(f"- {warning}" for warning in warnings)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    (REPORT_DIR / f"country-risk-{dataset['version']}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Import Damodaran country risk premium workbook")
    parser.add_argument("xlsx", help="Path to ctryprem*.xlsx")
    args = parser.parse_args()

    source = Path(args.xlsx)
    if not source.exists():
        raise FileNotFoundError(source)

    wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    dataset = {
        **find_metadata(wb),
        "sourceFile": source.name,
        "countries": extract_countries(wb),
    }
    warnings = validate_dataset(dataset)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / f"{dataset['version']}.json").write_text(
        json.dumps(dataset, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    write_latest_ts(dataset)
    write_report(dataset, warnings)
    print(f"Imported {len(dataset['countries'])} countries for {dataset['version']}")


if __name__ == "__main__":
    main()
